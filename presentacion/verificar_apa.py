# -*- coding: utf-8 -*-
from docx import Document

doc = Document(r'C:\Users\USER\Pictures\Nueva carpeta\RutaQuilla_ProyectoFinal_APA7.docx')

print('=== VERIFICACION APA 7 ===')
print()

# 1. Margins
s = doc.sections[0]
print('1. MARGENES:')
print(f'   Superior: {s.top_margin / 914400:.2f} in (esperado: 1.00)')
print(f'   Inferior: {s.bottom_margin / 914400:.2f} in (esperado: 1.00)')
print(f'   Izquierdo: {s.left_margin / 914400:.2f} in (esperado: 1.00)')
print(f'   Derecho: {s.right_margin / 914400:.2f} in (esperado: 1.00)')

# 2. Font
print('\n2. FUENTE BASE:')
style = doc.styles['Normal']
print(f'   Nombre: {style.font.name} (esperado: Times New Roman)')
print(f'   Tamano: {style.font.size / 12700}pt (esperado: 12)')

# 3. Title page
print('\n3. PORTADA (primeros 15 parrafos):')
align_map = {0: 'LEFT', 1: 'CENTER', 2: 'RIGHT', 3: 'JUSTIFY'}
for i in range(min(15, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    a = p.paragraph_format.alignment
    a_str = align_map.get(a, str(a))
    bold = any(r.bold for r in p.runs) if p.runs else False
    text = p.text[:70] if p.text else '(vacio)'
    print(f'   P{i}: "{text}" | align={a_str} | bold={bold}')

# 4. Headings sample
print('\n4. ENCABEZADOS (primeros encontrados):')
count = 0
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() and any(r.bold for r in p.runs if r.text.strip()):
        a = p.paragraph_format.alignment
        a_str = align_map.get(a, str(a))
        fi = p.paragraph_format.first_line_indent
        print(f'   P{i}: "{p.text[:60]}" | align={a_str} | indent={fi}')
        count += 1
        if count >= 8:
            break

# 5. Body paragraph sample
print('\n5. PARRAFOS DE CUERPO (muestra):')
count = 0
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() and not any(r.bold for r in p.runs if r.text.strip()):
        fi = p.paragraph_format.first_line_indent
        fi_val = f'{fi / 914400:.2f}in' if fi else 'None'
        ls = p.paragraph_format.line_spacing_rule
        font_name = p.runs[0].font.name if p.runs else 'N/A'
        print(f'   P{i}: "{p.text[:55]}..." | indent={fi_val} | spacing={ls} | font={font_name}')
        count += 1
        if count >= 5:
            break

# 6. Tables
print(f'\n6. TABLAS: {len(doc.tables)} tablas')

# 7. References
print('\n7. REFERENCIAS:')
for i, p in enumerate(doc.paragraphs):
    if 'Referencias' in p.text or 'Alcaldia' in p.text or 'MongoDB' in p.text:
        a = p.paragraph_format.alignment
        a_str = align_map.get(a, str(a))
        li = p.paragraph_format.left_indent
        li_val = f'{li / 914400:.2f}in' if li else 'None'
        print(f'   P{i}: "{p.text[:65]}" | align={a_str} | left_indent={li_val}')

# 8. Header
print('\n8. ENCABEZADO (num. pagina):')
for section in doc.sections:
    for p in section.header.paragraphs:
        a = p.alignment
        a_str = align_map.get(a, str(a))
        print(f'   Header align={a_str} | fields={len(p.runs)} runs')

print('\n=== VERIFICACION COMPLETADA ===')
