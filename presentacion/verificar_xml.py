# -*- coding: utf-8 -*-
from docx import Document
from lxml import etree

doc = Document(r'C:\Users\USER\Pictures\Nueva carpeta\RutaQuilla_ProyectoFinal_APA7.docx')
ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

print('=== PORTADA ===')
for i in range(12):
    p = doc.paragraphs[i]._element
    ppr = p.find(f'{{{ns}}}pPr')
    jc = ppr.find(f'{{{ns}}}jc') if ppr is not None else None
    jc_val = jc.get(f'{{{ns}}}val') if jc is not None else 'NONE'
    
    # Check bold in run
    runs = p.findall(f'{{{ns}}}r')
    bold = False
    text = ''
    for r in runs:
        t = r.find(f'{{{ns}}}t')
        if t is not None and t.text:
            text = t.text[:50]
        rpr = r.find(f'{{{ns}}}rPr')
        if rpr is not None:
            b = rpr.find(f'{{{ns}}}b')
            if b is not None:
                bval = b.get(f'{{{ns}}}val')
                if bval is None or bval == 'true' or bval == '1':
                    bold = True
    print(f'P{i}: jc={jc_val} | bold={bold} | "{text}"')

print('\n=== ENCABEZADOS (muestra) ===')
for i in range(12, min(50, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip() and any(r.bold for r in p.runs):
        ppr_el = p._element.find(f'{{{ns}}}pPr')
        jc = ppr_el.find(f'{{{ns}}}jc') if ppr_el is not None else None
        jc_val = jc.get(f'{{{ns}}}val') if jc is not None else 'NONE'
        ind = ppr_el.find(f'{{{ns}}}ind') if ppr_el is not None else None
        fl = ind.get(f'{{{ns}}}firstLine') if ind is not None else 'NONE'
        print(f'P{i}: jc={jc_val} | indent={fl} | "{p.text[:55]}"')

print('\n=== PARRAFO CUERPO (muestra) ===')
count = 0
for i in range(12, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.text.strip() and not any(r.bold for r in p.runs):
        ppr_el = p._element.find(f'{{{ns}}}pPr')
        if ppr_el is not None:
            ind = ppr_el.find(f'{{{ns}}}ind')
            fl = ind.get(f'{{{ns}}}firstLine') if ind is not None else 'NONE'
            sp = ppr_el.find(f'{{{ns}}}spacing')
            line = sp.get(f'{{{ns}}}line') if sp is not None else 'NONE'
        else:
            fl = 'NO pPr'
            line = 'NO pPr'
        fn = p.runs[0].font.name if p.runs else 'N/A'
        print(f'P{i}: indent={fl} | spacing={line} | font={fn} | "{p.text[:45]}"')
        count += 1
        if count >= 5:
            break

print('\n=== REFERENCIAS ===')
for i in range(len(doc.paragraphs) - 10, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.text.strip():
        ppr_el = p._element.find(f'{{{ns}}}pPr')
        jc = ppr_el.find(f'{{{ns}}}jc') if ppr_el is not None else None
        jc_val = jc.get(f'{{{ns}}}val') if jc is not None else 'NONE'
        ind = ppr_el.find(f'{{{ns}}}ind') if ppr_el is not None else None
        hang = ind.get(f'{{{ns}}}hanging') if ind is not None else 'NONE'
        left = ind.get(f'{{{ns}}}left') if ind is not None else 'NONE'
        print(f'P{i}: jc={jc_val} | left={left} | hanging={hang} | "{p.text[:50]}"')

print('\nVerificacion completada!')
